"""File text extraction for TXT and DOCX (and unsupported-type handling)."""
from io import BytesIO

import pytest

from scholar.extract import extract_text


def test_txt_extraction():
    assert extract_text("cv.txt", b"Hello\nworld") == "Hello\nworld"


def test_md_extraction():
    assert extract_text("notes.md", b"# Title\nbody") == "# Title\nbody"


def test_docx_extraction():
    import docx

    doc = docx.Document()
    doc.add_paragraph("Ayesha Khan")
    doc.add_paragraph("BSc Computer Science, GPA 3.8")
    buf = BytesIO()
    doc.save(buf)

    text = extract_text("cv.docx", buf.getvalue())
    assert "Ayesha Khan" in text
    assert "GPA 3.8" in text


def test_unsupported_type_raises():
    with pytest.raises(ValueError):
        extract_text("photo.png", b"\x89PNG")
